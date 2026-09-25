#!/usr/bin/env python3
"""Compare HELIOS genetic ancestry with recorded race and ethnic heritage.

The script reads controlled participant-level inputs but writes aggregate
tables and non-identifiable figures only. No participant identifiers,
individual PCA coordinates, or individual ancestry records are exported.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import platform
from pathlib import Path

import numpy as np
import pandas as pd
from PIL import Image, ImageDraw, ImageFont


GROUPS = ["chinese", "indian", "malay"]
GENETIC_ORDER = [
    "chinese",
    "indian",
    "malay",
    "chinese-malay",
    "indian-chinese",
    "malay-indian",
    "other",
]
GENETIC_MEMBERS = {
    "chinese": {"chinese"},
    "indian": {"indian"},
    "malay": {"malay"},
    "chinese-malay": {"chinese", "malay"},
    "indian-chinese": {"indian", "chinese"},
    "malay-indian": {"malay", "indian"},
    "other": {"other"},
}
DISPLAY = {
    "chinese": "Chinese",
    "indian": "Indian",
    "malay": "Malay",
    "chinese-malay": "Chinese–Malay",
    "indian-chinese": "Chinese–Indian",
    "malay-indian": "Indian–Malay",
    "other": "Other",
    "missing": "Not recorded",
}
COLOURS = {
    "chinese": "#1B9E77",
    "indian": "#D95F02",
    "malay": "#7570B3",
    "chinese-malay": "#2C7FB8",
    "indian-chinese": "#C44E52",
    "malay-indian": "#E6AB02",
    "other": "#6B7280",
    "missing": "#D1D5DB",
}
PC_COLUMNS = [f"PC{i}" for i in range(1, 6)]
Z95 = 1.959963984540054


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--annotation", type=Path, required=True)
    parser.add_argument("--core", type=Path, required=True)
    parser.add_argument("--interview", type=Path, required=True)
    parser.add_argument("--data-dictionary", type=Path)
    parser.add_argument("--full22", type=Path)
    parser.add_argument("--covariates10", type=Path)
    parser.add_argument("--sg100k-covariates", type=Path)
    parser.add_argument("--gwas22-chinese", type=Path)
    parser.add_argument("--gwas22-indian", type=Path)
    parser.add_argument("--gwas22-malay", type=Path)
    parser.add_argument("--output-root", type=Path, required=True)
    parser.add_argument("--minimum-cell-size", type=int, default=5)
    return parser.parse_args()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def write_tsv(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    frame.to_csv(path, sep="\t", index=False, na_rep="")


def write_text(text: str, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text.rstrip() + "\n", encoding="utf-8")


def parse_annotation(path: Path) -> pd.DataFrame:
    """Parse the SG100K annotation's unquoted JSON-valued CSV columns."""
    decoder = json.JSONDecoder()
    records = []
    with path.open(encoding="utf-8") as handle:
        header = next(handle).rstrip("\n").split(",")
        expected = [
            "npmid",
            "ilmn_meta",
            "sg100k_meta",
            "sample_filters",
            "sample_flags",
            "grids_meta",
            "sample_qc",
        ]
        if header != expected:
            raise ValueError(f"Unexpected annotation header: {header}")
        for line_number, line in enumerate(handle, 2):
            npmid, payload = line.rstrip("\n").split(",", 1)
            values = []
            position = 0
            try:
                for _ in range(6):
                    while position < len(payload) and payload[position] in " \t,":
                        position += 1
                    value, position = decoder.raw_decode(payload, position)
                    values.append(value)
            except Exception as exc:
                raise ValueError(
                    f"Could not parse annotation line {line_number}: {exc}"
                ) from exc
            _, meta, sample_filters, sample_flags, grids, _ = values
            scores = grids.get("pca_scores") or []
            records.append(
                {
                    "npmid": npmid,
                    "participant_id": meta.get("participant_id"),
                    "administrative_race": meta.get("race"),
                    "cohort": meta.get("cohort"),
                    "genetic_ancestry": grids.get("ancestry"),
                    "sample_filters": sample_filters,
                    "sample_flags": sample_flags,
                    **{
                        f"PC{i + 1}": scores[i] if i < len(scores) else np.nan
                        for i in range(5)
                    },
                }
            )
    frame = pd.DataFrame.from_records(records)
    if frame["npmid"].duplicated().any():
        raise ValueError("Annotation contains duplicate NPM IDs")
    if frame["participant_id"].duplicated().any():
        raise ValueError("Annotation contains duplicate participant IDs")
    return frame


def normalise_group(value: object) -> str | None:
    if pd.isna(value):
        return None
    text = str(value).strip().lower()
    mapping = {
        "chinese": "chinese",
        "indian": "indian",
        "malay": "malay",
        "other": "other",
        "others": "other",
    }
    return mapping.get(text, text or None)


def collapse_consistent(
    frame: pd.DataFrame, id_column: str, value_columns: list[str]
) -> tuple[pd.DataFrame, int, int]:
    """Collapse repeated records and exclude identifiers with conflicting values."""
    duplicate_rows = int(frame.duplicated(id_column, keep=False).sum())
    conflict_ids = []
    for identifier, group in frame.groupby(id_column, dropna=False, sort=False):
        conflict = False
        for column in value_columns:
            values = group[column].dropna().astype(str).str.strip().unique()
            if len(values) > 1:
                conflict = True
                break
        if conflict:
            conflict_ids.append(identifier)
    clean = frame.loc[
        ~frame[id_column].isin(conflict_ids), [id_column] + value_columns
    ].copy()
    clean = clean.drop_duplicates([id_column] + value_columns).drop_duplicates(
        id_column
    )
    return clean, duplicate_rows, len(conflict_ids)


def wilson(successes: int, total: int) -> tuple[float, float]:
    if total == 0:
        return np.nan, np.nan
    proportion = successes / total
    denominator = 1 + Z95**2 / total
    centre = (proportion + Z95**2 / (2 * total)) / denominator
    margin = (
        Z95
        * math.sqrt(
            proportion * (1 - proportion) / total + Z95**2 / (4 * total**2)
        )
        / denominator
    )
    return max(0.0, centre - margin), min(1.0, centre + margin)


def category_for(recorded_group: str, genetic_group: str) -> str:
    if genetic_group == recorded_group:
        return "exact_core_match"
    members = GENETIC_MEMBERS.get(genetic_group, set())
    if recorded_group in members and len(members) == 2:
        return "intermediate_involving_group"
    return "genetic_label_not_involving_group"


def classification_summary(
    frame: pd.DataFrame, recorded_column: str, source_label: str
) -> pd.DataFrame:
    rows = []
    for group in GROUPS:
        subset = frame.loc[
            frame[recorded_column].eq(group) & frame["genetic_ancestry"].notna()
        ]
        total = len(subset)
        categories = subset.apply(
            lambda row: category_for(group, row["genetic_ancestry"]), axis=1
        ).value_counts()
        row = {
            "source": source_label,
            "recorded_group": DISPLAY[group],
            "total": total,
        }
        for category in [
            "exact_core_match",
            "intermediate_involving_group",
            "genetic_label_not_involving_group",
        ]:
            count = int(categories.get(category, 0))
            low, high = wilson(count, total)
            row[f"{category}_n"] = count
            row[f"{category}_percentage"] = (
                100 * count / total if total else np.nan
            )
            row[f"{category}_ci_low"] = 100 * low
            row[f"{category}_ci_high"] = 100 * high
        rows.append(row)
    return pd.DataFrame(rows)


def suppress_crosstab(
    frame: pd.DataFrame,
    row_column: str,
    column_column: str,
    minimum_cell_size: int,
) -> pd.DataFrame:
    table = (
        frame.groupby([row_column, column_column], dropna=False)
        .size()
        .rename("raw_count")
        .reset_index()
    )
    totals = table.groupby(row_column)["raw_count"].transform("sum")
    table["row_percentage_raw"] = 100 * table["raw_count"] / totals
    table["suppressed"] = table["raw_count"].between(1, minimum_cell_size - 1)
    for _, indices in table.groupby(row_column).groups.items():
        indices = list(indices)
        primary = [i for i in indices if bool(table.loc[i, "suppressed"])]
        if len(primary) == 1:
            candidates = [
                i
                for i in indices
                if int(table.loc[i, "raw_count"]) >= minimum_cell_size
            ]
            if candidates:
                companion = min(
                    candidates, key=lambda i: (table.loc[i, "raw_count"], i)
                )
                table.loc[companion, "suppressed"] = True
    table["count"] = table["raw_count"].astype(str)
    table.loc[table["suppressed"], "count"] = f"<{minimum_cell_size}"
    table["row_percentage"] = table["row_percentage_raw"]
    table.loc[table["suppressed"], "row_percentage"] = np.nan
    return table[
        [row_column, column_column, "count", "row_percentage", "suppressed"]
    ]


def chi_square_three_by_two(
    events: dict[str, tuple[int, int]]
) -> tuple[float, float]:
    observed = np.array(
        [
            [events[group][0], events[group][1] - events[group][0]]
            for group in GROUPS
        ],
        dtype=float,
    )
    expected = observed.sum(axis=1, keepdims=True) @ (
        observed.sum(axis=0, keepdims=True) / observed.sum()
    )
    statistic = float(np.sum((observed - expected) ** 2 / expected))
    p_value = math.exp(-statistic / 2)
    return statistic, p_value


def log_choose(n: int, k: int) -> float:
    if k < 0 or k > n:
        return -math.inf
    return (
        math.lgamma(n + 1)
        - math.lgamma(k + 1)
        - math.lgamma(n - k + 1)
    )


def fisher_two_sided(a: int, b: int, c: int, d: int) -> float:
    row1, row2 = a + b, c + d
    successes = a + c
    total = row1 + row2
    lower = max(0, row1 - (total - successes))
    upper = min(row1, successes)

    def log_probability(x: int) -> float:
        return (
            log_choose(successes, x)
            + log_choose(total - successes, row1 - x)
            - log_choose(total, row1)
        )

    observed = log_probability(a)
    selected = []
    for x in range(lower, upper + 1):
        value = log_probability(x)
        if value <= observed + 1e-12:
            selected.append(value)
    peak = max(selected)
    return min(
        1.0, math.exp(peak) * sum(math.exp(value - peak) for value in selected)
    )


def pairwise_effects(a: int, n1: int, c: int, n2: int) -> dict[str, float]:
    b, d = n1 - a, n2 - c
    corrected = [float(a), float(b), float(c), float(d)]
    if min(corrected) == 0:
        corrected = [value + 0.5 for value in corrected]
    aa, bb, cc, dd = corrected
    risk_ratio = (aa / (aa + bb)) / (cc / (cc + dd))
    rr_se = math.sqrt(
        1 / aa - 1 / (aa + bb) + 1 / cc - 1 / (cc + dd)
    )
    odds_ratio = aa * dd / (bb * cc)
    or_se = math.sqrt(1 / aa + 1 / bb + 1 / cc + 1 / dd)
    return {
        "risk_ratio": risk_ratio,
        "risk_ratio_ci_low": math.exp(math.log(risk_ratio) - Z95 * rr_se),
        "risk_ratio_ci_high": math.exp(math.log(risk_ratio) + Z95 * rr_se),
        "odds_ratio": odds_ratio,
        "odds_ratio_ci_low": math.exp(math.log(odds_ratio) - Z95 * or_se),
        "odds_ratio_ci_high": math.exp(math.log(odds_ratio) + Z95 * or_se),
        "fisher_exact_p": fisher_two_sided(a, b, c, d),
    }


def statistical_tests(summary: pd.DataFrame) -> pd.DataFrame:
    endpoints = {
        "not_exact_core_match": lambda row: row["total"]
        - row["exact_core_match_n"],
        "intermediate_involving_group": lambda row: row[
            "intermediate_involving_group_n"
        ],
        "genetic_label_not_involving_group": lambda row: row[
            "genetic_label_not_involving_group_n"
        ],
    }
    rows = []
    indexed = summary.set_index(summary["recorded_group"].str.lower())
    for endpoint, getter in endpoints.items():
        events = {
            group: (
                int(getter(indexed.loc[group])),
                int(indexed.loc[group, "total"]),
            )
            for group in GROUPS
        }
        chi2, p_value = chi_square_three_by_two(events)
        rows.append(
            {
                "endpoint": endpoint,
                "comparison": "Chinese vs Indian vs Malay",
                "test": "Pearson chi-square",
                "statistic": chi2,
                "degrees_of_freedom": 2,
                "p_value": p_value,
            }
        )
        for second in ["indian", "malay"]:
            a, n1 = events["chinese"]
            c, n2 = events[second]
            effects = pairwise_effects(a, n1, c, n2)
            rows.append(
                {
                    "endpoint": endpoint,
                    "comparison": f"Chinese vs {DISPLAY[second]}",
                    "test": "two-sided Fisher exact",
                    "chinese_events": a,
                    "chinese_total": n1,
                    "comparison_events": c,
                    "comparison_total": n2,
                    **effects,
                }
            )
    return pd.DataFrame(rows)


def parental_pair(row: pd.Series) -> str | None:
    father = normalise_group(row.get("father_ethnic_heritage"))
    mother = normalise_group(row.get("mother_ethnic_heritage"))
    if father is None or mother is None:
        return None
    order = {"chinese": 0, "indian": 1, "malay": 2, "other": 3}
    pair = sorted([father, mother], key=lambda item: order.get(item, 9))
    return "-".join(pair)


def parental_validation(frame: pd.DataFrame) -> pd.DataFrame:
    expected = {
        "chinese": "chinese-chinese",
        "indian": "indian-indian",
        "malay": "malay-malay",
        "chinese-malay": "chinese-malay",
        "indian-chinese": "chinese-indian",
        "malay-indian": "indian-malay",
        "other": "other-other",
    }
    rows = []
    for genetic in GENETIC_ORDER:
        subset = frame.loc[frame["genetic_ancestry"].eq(genetic)]
        known = subset["parental_pair"].notna()
        denominator = int(known.sum())
        matches = int(
            subset.loc[known, "parental_pair"].eq(expected[genetic]).sum()
        )
        low, high = wilson(matches, denominator)
        rows.append(
            {
                "genetic_ancestry": DISPLAY[genetic],
                "total_n": len(subset),
                "both_parental_heritage_recorded_n": denominator,
                "expected_parental_pair": expected[genetic]
                .replace("-", "–")
                .title(),
                "matching_parental_pair_n": matches,
                "matching_parental_pair_percentage": (
                    100 * matches / denominator if denominator else np.nan
                ),
                "ci_low": 100 * low,
                "ci_high": 100 * high,
            }
        )
    return pd.DataFrame(rows)


def relatedness_summary(frame: pd.DataFrame) -> pd.DataFrame:
    def has_relatedness_flag(flags: object) -> bool:
        if not isinstance(flags, list):
            return False
        return any(
            "degree" in str(value).lower()
            or "duplicate" in str(value).lower()
            for value in flags
        )

    flagged = frame["sample_flags"].map(has_relatedness_flag)
    rows = []
    for genetic in GENETIC_ORDER:
        mask = frame["genetic_ancestry"].eq(genetic)
        total = int(mask.sum())
        count = int(flagged.loc[mask].sum())
        low, high = wilson(count, total)
        rows.append(
            {
                "genetic_ancestry": DISPLAY[genetic],
                "total_n": total,
                "relatedness_or_duplicate_flag_n": count,
                "percentage": 100 * count / total if total else np.nan,
                "ci_low": 100 * low,
                "ci_high": 100 * high,
            }
        )
    return pd.DataFrame(rows)


def centroid_sensitivity(frame: pd.DataFrame) -> pd.DataFrame:
    core = frame["genetic_ancestry"].isin(GROUPS)
    values = frame[PC_COLUMNS].to_numpy(dtype=float)
    if not np.isfinite(values).all():
        raise ValueError("GRID PCA scores are incomplete")
    schemes = [
        ("raw_PC1_PC2", 2, False),
        ("raw_PC1_PC3", 3, False),
        ("raw_PC1_PC5", 5, False),
        ("standardised_PC1_PC5", 5, True),
    ]
    rows = []
    for scheme, dimensions, standardised in schemes:
        matrix = values[:, :dimensions].copy()
        if standardised:
            centre = matrix[core].mean(axis=0)
            scale = matrix[core].std(axis=0, ddof=1)
            matrix = (matrix - centre) / scale
        centroids = np.vstack(
            [
                matrix[frame["genetic_ancestry"].eq(group)].mean(axis=0)
                for group in GROUPS
            ]
        )
        distance = (
            (matrix[:, None, :] - centroids[None, :, :]) ** 2
        ).sum(axis=2)
        nearest = np.array(GROUPS, dtype=object)[np.argmin(distance, axis=1)]
        for genetic in GENETIC_ORDER:
            mask = frame["genetic_ancestry"].eq(genetic).to_numpy()
            total = int(mask.sum())
            for target in GROUPS:
                count = int(np.sum(nearest[mask] == target))
                rows.append(
                    {
                        "sensitivity": scheme,
                        "genetic_ancestry": DISPLAY[genetic],
                        "nearest_core_centroid": DISPLAY[target],
                        "n": count,
                        "total_n": total,
                        "percentage": (
                            100 * count / total if total else np.nan
                        ),
                    }
                )
    return pd.DataFrame(rows)


def font(size: int, bold: bool = False) -> ImageFont.ImageFont:
    candidates = [
        "/System/Library/Fonts/HelveticaNeue.ttc",
        "/System/Library/Fonts/Helvetica.ttc",
        (
            "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
            if bold
            else "/System/Library/Fonts/Supplemental/Arial.ttf"
        ),
    ]
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size)
        except OSError:
            continue
    return ImageFont.load_default()


def rgb(hex_colour: str) -> tuple[int, int, int]:
    text = hex_colour.lstrip("#")
    return tuple(int(text[i : i + 2], 16) for i in (0, 2, 4))


def blend(hex_colour: str, alpha: float) -> tuple[int, int, int]:
    base = rgb(hex_colour)
    return tuple(
        round(255 - alpha * (255 - value)) for value in base
    )


def save_figure(image: Image.Image, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, dpi=(300, 300))
    image.save(path.with_suffix(".pdf"), "PDF", resolution=300)


def scatter_plot(
    frame: pd.DataFrame,
    category_column: str,
    title: str,
    subtitle: str,
    path: Path,
    category_order: list[str],
    highlight_only: bool = False,
) -> None:
    width, height = 1650, 1200
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    left, top, right, bottom = (145, 135, 1260, 1050)
    draw.text((70, 38), title, fill="#17202A", font=font(34, True))
    draw.text((72, 82), subtitle, fill="#5B6573", font=font(18))
    x = frame["PC1"].to_numpy(dtype=float)
    y = frame["PC2"].to_numpy(dtype=float)
    x_low, x_high = np.quantile(x, [0.002, 0.998])
    y_low, y_high = np.quantile(y, [0.002, 0.998])
    x_pad = 0.04 * (x_high - x_low)
    y_pad = 0.04 * (y_high - y_low)
    x_low, x_high = x_low - x_pad, x_high + x_pad
    y_low, y_high = y_low - y_pad, y_high + y_pad

    def sx(value: float) -> float:
        return left + (value - x_low) / (x_high - x_low) * (right - left)

    def sy(value: float) -> float:
        return bottom - (value - y_low) / (y_high - y_low) * (bottom - top)

    for fraction in np.linspace(0, 1, 6):
        px = left + fraction * (right - left)
        py = bottom - fraction * (bottom - top)
        draw.line((px, top, px, bottom), fill="#E5E7EB", width=1)
        draw.line((left, py, right, py), fill="#E5E7EB", width=1)
        xv = x_low + fraction * (x_high - x_low)
        yv = y_low + fraction * (y_high - y_low)
        draw.text(
            (px, bottom + 12),
            f"{xv:.2f}",
            anchor="ma",
            fill="#5B6573",
            font=font(14),
        )
        draw.text(
            (left - 12, py),
            f"{yv:.2f}",
            anchor="rm",
            fill="#5B6573",
            font=font(14),
        )
    draw.line((left, bottom, right, bottom), fill="#17202A", width=2)
    draw.line((left, top, left, bottom), fill="#17202A", width=2)
    draw.text(
        ((left + right) / 2, bottom + 58),
        "GRID PC1",
        anchor="mm",
        fill="#17202A",
        font=font(20),
    )
    draw.text(
        (50, (top + bottom) / 2),
        "GRID PC2",
        anchor="mm",
        fill="#17202A",
        font=font(20),
    )

    categories = (
        frame[category_column].fillna("missing").astype(str).to_numpy()
    )
    if highlight_only:
        for xv, yv in zip(x, y):
            px, py = sx(xv), sy(yv)
            if left <= px <= right and top <= py <= bottom:
                draw.ellipse(
                    (px - 1, py - 1, px + 1, py + 1), fill="#D1D5DB"
                )
    for category in category_order:
        mask = categories == category
        colour = COLOURS.get(category, "#6B7280")
        alpha = 0.72 if highlight_only else (
            0.50 if category != "missing" else 0.35
        )
        radius = 3 if highlight_only else 2
        for xv, yv in zip(x[mask], y[mask]):
            px, py = sx(xv), sy(yv)
            if left <= px <= right and top <= py <= bottom:
                draw.ellipse(
                    (
                        px - radius,
                        py - radius,
                        px + radius,
                        py + radius,
                    ),
                    fill=blend(colour, alpha),
                )
    legend_x, legend_y = 1305, 165
    legend_index = 0
    for category in category_order:
        count = int(np.sum(categories == category))
        if count == 0:
            continue
        y0 = legend_y + legend_index * 52
        legend_index += 1
        draw.ellipse(
            (legend_x, y0, legend_x + 18, y0 + 18),
            fill=COLOURS.get(category, "#6B7280"),
        )
        label = DISPLAY.get(
            category, category.replace("_", " ").title()
        )
        draw.text(
            (legend_x + 30, y0 + 9),
            f"{label} (N={count:,})",
            anchor="lm",
            fill="#17202A",
            font=font(17),
        )
    save_figure(image, path)


def composition_plot(
    summary: pd.DataFrame, title: str, path: Path
) -> None:
    width, height = 1450, 780
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 38), title, fill="#17202A", font=font(32, True))
    draw.text(
        (72, 82),
        "Exact core match, pairwise intermediate involving the recorded group, and a genetic label not involving that group",
        fill="#5B6573",
        font=font(17),
    )
    left, right = 300, 1360
    top, bottom = 165, 640
    for fraction in np.linspace(0, 1, 6):
        x = left + fraction * (right - left)
        draw.line((x, top, x, bottom), fill="#E5E7EB", width=1)
        draw.text(
            (x, bottom + 18),
            f"{fraction * 100:.0f}%",
            anchor="ma",
            fill="#5B6573",
            font=font(15),
        )
    keys = [
        ("exact_core_match_percentage", "#4C9F70", "Exact core match"),
        (
            "intermediate_involving_group_percentage",
            "#E6AB02",
            "Pairwise intermediate",
        ),
        (
            "genetic_label_not_involving_group_percentage",
            "#C44E52",
            "Label not involving group",
        ),
    ]
    for row_index, (_, row) in enumerate(summary.iterrows()):
        y0 = top + 70 + row_index * 125
        draw.text(
            (left - 22, y0 + 24),
            row["recorded_group"],
            anchor="rm",
            fill="#17202A",
            font=font(20, True),
        )
        cursor = left
        for key, colour, _ in keys:
            value = float(row[key])
            end = cursor + value / 100 * (right - left)
            draw.rectangle((cursor, y0, end, y0 + 48), fill=colour)
            if value >= 4:
                draw.text(
                    ((cursor + end) / 2, y0 + 24),
                    f"{value:.1f}%",
                    anchor="mm",
                    fill="white",
                    font=font(15, True),
                )
            cursor = end
    legend_x, legend_y = 320, 700
    for index, (_, colour, label) in enumerate(keys):
        x = legend_x + index * 355
        draw.rectangle((x, legend_y, x + 22, legend_y + 22), fill=colour)
        draw.text(
            (x + 32, legend_y + 11),
            label,
            anchor="lm",
            fill="#17202A",
            font=font(15),
        )
    save_figure(image, path)


def variable_definitions(
    data_dictionary: Path | None,
) -> pd.DataFrame:
    variables = [
        "FREG5_Race",
        "FREG6_Race_Other",
        "FIAQ10_Demo1",
        "FIAQ14_1_Demo6",
        "FIAQ14_2_Demo7",
    ]
    if data_dictionary is None:
        return pd.DataFrame(
            {
                "variable": variables,
                "definition": ["not checked"] * len(variables),
            }
        )
    from openpyxl import load_workbook

    workbook = load_workbook(
        data_dictionary, read_only=True, data_only=True
    )
    records = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            for variable in variables:
                if variable in values:
                    records.append(
                        {
                            "variable": variable,
                            "definition": (
                                values[5] if len(values) > 5 else ""
                            ),
                            "allowed_values": (
                                values[10] if len(values) > 10 else ""
                            ),
                            "source_sheet": worksheet.title,
                        }
                    )
    result = pd.DataFrame(records).drop_duplicates("variable")
    missing = set(variables) - set(result.get("variable", []))
    if missing:
        raise ValueError(f"Definitions not found for: {sorted(missing)}")
    return result


def analysis_group_checks(
    frame: pd.DataFrame, args: argparse.Namespace
) -> pd.DataFrame:
    rows = []
    if args.full22:
        full = pd.read_csv(
            args.full22, sep=r"\s+", dtype={"IID": str}
        )
        if full["IID"].duplicated().any():
            raise ValueError("The 22k analysis table contains duplicate IIDs")
        joined = frame[["npmid", "genetic_ancestry"]].merge(
            full[["IID", "Ancestry"]],
            left_on="npmid",
            right_on="IID",
            how="left",
        )
        for genetic in GENETIC_ORDER:
            subset = joined.loc[
                joined["genetic_ancestry"].eq(genetic)
            ]
            for label in ["Chinese", "Indian", "Malay"]:
                rows.append(
                    {
                        "analysis_source": (
                            "Combined 22k retained analysis table"
                        ),
                        "genetic_ancestry": DISPLAY[genetic],
                        "analysis_group": label,
                        "unique_participants_n": int(
                            subset["Ancestry"].eq(label).sum()
                        ),
                    }
                )
            rows.append(
                {
                    "analysis_source": (
                        "Combined 22k retained analysis table"
                    ),
                    "genetic_ancestry": DISPLAY[genetic],
                    "analysis_group": "No ancestry-specific label",
                    "unique_participants_n": int(
                        subset["Ancestry"].isna().sum()
                    ),
                }
            )
    if args.covariates10:
        cov10 = pd.read_csv(
            args.covariates10, sep=r"\s+", dtype={"IID": str}
        )
        reduced = cov10[["IID", "Ancestry"]].drop_duplicates()
        if reduced["IID"].duplicated().any():
            raise ValueError(
                "The 10k covariate table assigns an IID to multiple groups"
            )
        joined = frame[["npmid", "genetic_ancestry"]].merge(
            reduced, left_on="npmid", right_on="IID", how="left"
        )
        for genetic in GENETIC_ORDER:
            subset = joined.loc[
                joined["genetic_ancestry"].eq(genetic)
            ]
            for label in ["Chinese", "Indian", "Malay"]:
                rows.append(
                    {
                        "analysis_source": (
                            "Retained 10k merged covariate table"
                        ),
                        "genetic_ancestry": DISPLAY[genetic],
                        "analysis_group": label,
                        "unique_participants_n": int(
                            subset["Ancestry"].eq(label).sum()
                        ),
                    }
                )
            rows.append(
                {
                    "analysis_source": (
                        "Retained 10k merged covariate table"
                    ),
                    "genetic_ancestry": DISPLAY[genetic],
                    "analysis_group": "Not present",
                    "unique_participants_n": int(
                        subset["Ancestry"].isna().sum()
                    ),
                }
            )
    gwas_paths = {
        "Current 22k Chinese covariate input": args.gwas22_chinese,
        "Current 22k Indian covariate input": args.gwas22_indian,
        "Current 22k Malay covariate input": args.gwas22_malay,
    }
    for label, path in gwas_paths.items():
        if path is None:
            continue
        input_frame = pd.read_csv(
            path, sep=r"\s+", usecols=["IID"], dtype={"IID": str}
        )
        ids = set(input_frame["IID"].dropna().astype(str))
        for genetic in GENETIC_ORDER:
            rows.append(
                {
                    "analysis_source": label,
                    "genetic_ancestry": DISPLAY[genetic],
                    "analysis_group": "Present in input",
                    "unique_participants_n": int(
                        frame.loc[
                            frame["genetic_ancestry"].eq(genetic),
                            "npmid",
                        ]
                        .isin(ids)
                        .sum()
                    ),
                }
            )
    return pd.DataFrame(rows)


def report_markdown(
    frame: pd.DataFrame,
    self_report: pd.DataFrame,
    asymmetry_tests: pd.DataFrame,
    parent: pd.DataFrame,
    availability: pd.DataFrame,
    group_checks: pd.DataFrame,
    centroid_results: pd.DataFrame,
) -> str:
    indexed = self_report.set_index("recorded_group")
    s_ch = indexed.loc["Chinese"]
    s_in = indexed.loc["Indian"]
    s_ml = indexed.loc["Malay"]
    p = parent.set_index("genetic_ancestry")
    strict_tests = asymmetry_tests.loc[
        asymmetry_tests["endpoint"].eq(
            "genetic_label_not_involving_group"
        )
    ].set_index("comparison")
    strict_global = strict_tests.loc["Chinese vs Indian vs Malay"]
    strict_indian = strict_tests.loc["Chinese vs Indian"]
    strict_malay = strict_tests.loc["Chinese vs Malay"]
    standardised_centroids = centroid_results.loc[
        centroid_results["sensitivity"].eq("standardised_PC1_PC5")
    ]
    core_centroid_rates = {}
    for group in ["Chinese", "Indian", "Malay"]:
        row = standardised_centroids.loc[
            standardised_centroids["genetic_ancestry"].eq(group)
            & standardised_centroids["nearest_core_centroid"].eq(group)
        ].iloc[0]
        core_centroid_rates[group] = float(row["percentage"])
    genetic_counts = frame["genetic_ancestry"].value_counts()
    self_n = int(frame["self_reported_ethnicity"].notna().sum())
    total = len(frame)
    availability_lines = []
    for _, row in availability.iterrows():
        availability_lines.append(
            f"- {row['administrative_race']}: "
            f"{int(row['self_reported_ethnicity_recorded_n']):,}/"
            f"{int(row['total_n']):,} ({row['percentage']:.1f}%)."
        )
    intermediate_assigned = None
    if not group_checks.empty:
        current = group_checks.loc[
            group_checks["analysis_source"].eq(
                "Combined 22k retained analysis table"
            )
        ]
        if not current.empty:
            intermediate_labels = {
                "Chinese–Malay",
                "Chinese–Indian",
                "Indian–Malay",
                "Other",
            }
            intermediate_assigned = int(
                current.loc[
                    current["genetic_ancestry"].isin(
                        intermediate_labels
                    )
                    & ~current["analysis_group"].eq(
                        "No ancestry-specific label"
                    ),
                    "unique_participants_n",
                ].sum()
            )
    gwas_sentence = ""
    if intermediate_assigned is not None:
        gwas_sentence = (
            "The combined 22k retained analysis table assigned ancestry-"
            "specific labels only to GRID core Chinese, Indian and Malay "
            f"participants. Intermediate/other participants assigned to an "
            f"ancestry-specific label: **{intermediate_assigned}**."
        )
    return f"""# HELIOS ancestry and self-reported ethnicity analysis

## Data sources

This analysis separates three variables previously described under the same ancestry label:

- grids_meta.ancestry: GRID genetic ancestry assignment and GRID PC1–PC5;
- FREG5_Race: race recorded from the NRIC; and
- FIAQ10_Demo1: participant-described ethnic heritage from the interviewer-administered questionnaire.

Parental ethnic heritage comes from FIAQ14_1_Demo6 and FIAQ14_2_Demo7. The analysis includes **{total:,}** genetically annotated HELIOS participants. Self-described ethnic heritage is available for **{self_n:,} ({100*self_n/total:.1f}%)**.

## Existing ancestry pipeline

The SG100K annotation provides seven genetic categories. The HELIOS counts are:

- Chinese: **{int(genetic_counts.get('chinese', 0)):,}**
- Indian: **{int(genetic_counts.get('indian', 0)):,}**
- Malay: **{int(genetic_counts.get('malay', 0)):,}**
- Chinese–Malay: **{int(genetic_counts.get('chinese-malay', 0)):,}**
- Chinese–Indian: **{int(genetic_counts.get('indian-chinese', 0)):,}**
- Indian–Malay: **{int(genetic_counts.get('malay-indian', 0)):,}**
- Other: **{int(genetic_counts.get('other', 0)):,}**

These supplied genetic categories, rather than a new PCA threshold, are the primary ancestry classification.

## Self-reported ethnicity variables

FREG5_Race should be described as **NRIC-recorded race**, not self-reported ancestry. FIAQ10_Demo1 is the appropriate self-described ethnic-heritage variable. Among participants present in both sources, the annotation race and FREG5_Race agree exactly after removal of one conflicting duplicate core record.

Questionnaire availability differs substantially by NRIC-recorded race:

{chr(10).join(availability_lines)}

Comparisons using FIAQ10_Demo1 therefore describe the questionnaire subset and may be affected by differential availability.

## Genetic ancestry classification

The older script that produced PCA_inferred.png copied the recorded label for every participant with a non-missing label and inferred only missing values. It is not an independent genetic-classification result and should not be used as evidence of concordance.

## Admixed/intermediate ancestry

The GRID annotation directly identifies Chinese–Malay, Chinese–Indian and Indian–Malay genetic categories. In participants with both parental ethnic-heritage variables recorded:

- Chinese–Indian: **{int(p.loc['Chinese–Indian', 'matching_parental_pair_n']):,}/{int(p.loc['Chinese–Indian', 'both_parental_heritage_recorded_n']):,} ({p.loc['Chinese–Indian', 'matching_parental_pair_percentage']:.1f}%)** reported one Chinese and one Indian parent.
- Chinese–Malay: **{int(p.loc['Chinese–Malay', 'matching_parental_pair_n']):,}/{int(p.loc['Chinese–Malay', 'both_parental_heritage_recorded_n']):,} ({p.loc['Chinese–Malay', 'matching_parental_pair_percentage']:.1f}%)** reported one Chinese and one Malay parent.
- Indian–Malay: **{int(p.loc['Indian–Malay', 'matching_parental_pair_n']):,}/{int(p.loc['Indian–Malay', 'both_parental_heritage_recorded_n']):,} ({p.loc['Indian–Malay', 'matching_parental_pair_percentage']:.1f}%)** reported one Indian and one Malay parent.

The Chinese–Indian category has the strongest direct support from parental heritage. The lower matching-parent proportions for Chinese–Malay and Indian–Malay support more cautious language: the genetic categories may reflect recent mixture, older admixture, continuous population structure, or limitations of categorical parental responses.

## Self-report/genetic concordance

Within the questionnaire subset:

- Chinese: exact core match **{s_ch['exact_core_match_percentage']:.2f}%**; pairwise intermediate involving Chinese **{s_ch['intermediate_involving_group_percentage']:.2f}%**; genetic label not involving Chinese **{s_ch['genetic_label_not_involving_group_percentage']:.2f}%**.
- Indian: exact core match **{s_in['exact_core_match_percentage']:.2f}%**; pairwise intermediate involving Indian **{s_in['intermediate_involving_group_percentage']:.2f}%**; genetic label not involving Indian **{s_in['genetic_label_not_involving_group_percentage']:.2f}%**.
- Malay: exact core match **{s_ml['exact_core_match_percentage']:.2f}%**; pairwise intermediate involving Malay **{s_ml['intermediate_involving_group_percentage']:.2f}%**; genetic label not involving Malay **{s_ml['genetic_label_not_involving_group_percentage']:.2f}%**.

## Direction of discordance

The apparent excess of Chinese discordance in the earlier PCA is not supported after using rates. Self-described Chinese participants have the **lowest**, not the highest, rate of a genetic label that does not involve their recorded group. The same ordering is present when NRIC-recorded race is used on the full annotated sample.

Rates of a strict non-involving genetic label differ across the three self-described groups (Pearson chi-square = **{strict_global['statistic']:.2f}**, df = 2, p = **{strict_global['p_value']:.2e}**). The risk among self-described Chinese participants is **{strict_indian['risk_ratio']:.3f} times** that among self-described Indian participants (95% CI {strict_indian['risk_ratio_ci_low']:.3f}–{strict_indian['risk_ratio_ci_high']:.3f}; Fisher p = {strict_indian['fisher_exact_p']:.2e}) and **{strict_malay['risk_ratio']:.3f} times** that among self-described Malay participants (95% CI {strict_malay['risk_ratio_ci_low']:.3f}–{strict_malay['risk_ratio_ci_high']:.3f}; Fisher p = {strict_malay['fisher_exact_p']:.2e}). These comparisons are descriptive of the questionnaire subset because ethnic-heritage availability is differential.

## Chinese–Indian intermediate ancestry

The GRID Chinese–Indian category contains **{int(genetic_counts.get('indian-chinese', 0)):,}** HELIOS participants. Its parental-heritage match is higher than for the other two pairwise genetic categories. This is compatible with a recent mixed-ancestry contribution in a subset, but the genetic label should not be interpreted as proof of parental background for every participant.

## Chinese–Malay intermediate ancestry

The GRID Chinese–Malay category contains **{int(genetic_counts.get('chinese-malay', 0)):,}** participants. Because Chinese and Malay populations are genetically closer and only a minority of participants with parental data report one Chinese and one Malay parent, describe this category as **Chinese–Malay intermediate genetic ancestry** unless additional ancestry-proportion evidence is available.

## Indian–Malay intermediate ancestry

The GRID Indian–Malay category contains **{int(genetic_counts.get('malay-indian', 0)):,}** participants. Within-Indian heterogeneity and incomplete parental data remain relevant limitations.

## GWAS inclusion consequences

{gwas_sentence}

The retained 10k merged covariate table also maps its Chinese, Indian and Malay Ancestry values exclusively to the corresponding GRID core category. Thus the retained analysis Ancestry variable is a genetic-analysis grouping, not a self-report variable.

The pooled 22k phenotype PCA nevertheless contained 826 pairwise intermediate or Other GRID participants (3.68% of N=22,422). Because the retained phenotype regression used only Indian and Malay indicators, all 826 entered its Chinese reference coding; seven core GRID Indian or Malay participants were coded the same way. This affected phenotype construction but not membership of the ancestry-specific GWASs. A direct seven-category GRID sensitivity and a core-only refit showed that the coding had a negligible effect on PC1.

## Sensitivity analyses

- Nearest-centroid descriptions were repeated using PC1–PC2, PC1–PC3, PC1–PC5 and standardised PC1–PC5. With standardised PC1–PC5, the original core group remained the nearest centroid for **{core_centroid_rates['Chinese']:.1f}%** of core Chinese, **{core_centroid_rates['Indian']:.1f}%** of core Indian and **{core_centroid_rates['Malay']:.1f}%** of core Malay participants. Allocation of intermediate groups changed with PC scaling, confirming that this descriptive procedure cannot reconstruct the original GRID classification rule.
- Relatedness/duplicate flags were summarised by genetic category. These flags do not establish whether a visible bridge is family-driven; a kinship matrix would be required for that test.
- The available SG100K covariate file contains subcohort values for core groups but not the pairwise intermediate categories, so a valid intermediate-group batch/subcohort comparison could not be completed from this file.
- Replacing the original phenotype ancestry dummies with all seven GRID categories produced loading congruence 0.99999994 and score correlation r=0.999851. Restricting phenotype construction to the core GRID groups produced congruence 0.99998224 and r=0.999882.
- Within the same 13,102-person questionnaire subset, adding participant-described ethnic heritage produced loading congruence 0.99999997 and r=0.999729. Excluding discordant records produced congruence 0.99999884 and r=0.999982. Full results are in [`phenotype_sensitivity`](phenotype_sensitivity/PHENOTYPE_ETHNICITY_SENSITIVITY.md).

## Interpretation

Self-described ethnic heritage, NRIC-recorded race and genetic ancestry are related but non-equivalent variables. The major genetic clusters align closely with both recorded variables, while GRID also identifies pairwise intermediate structure. The data do not support the hypothesis that self-described Chinese participants are disproportionately discordant after accounting for group size. Correcting the pooled phenotype ancestry coding and modelling self-described ethnicity separately did not materially change the cognitive factor.

## Limitations

- Self-described ethnic heritage is missing non-randomly across NRIC race groups.
- GRID ancestry proportions and the exact classification thresholds were not present in the retained local archive.
- Parental ethnic heritage uses broad categories and cannot distinguish all forms or generations of admixture.
- PCA and categorical ancestry assignments cannot identify why a participant uses a particular ethnic label.
- A formal batch analysis and family-cluster analysis require source variables not present for the intermediate groups in the retained files.
"""


def main() -> None:
    args = parse_args()
    root = args.output_root
    tables = root / "tables"
    figures = root / "figures"
    tables.mkdir(parents=True, exist_ok=True)
    figures.mkdir(parents=True, exist_ok=True)

    annotation_all = parse_annotation(args.annotation)
    helios = annotation_all.loc[
        annotation_all["cohort"].astype(str).str.upper().eq("HELIOS")
    ].copy()
    helios["administrative_race"] = helios[
        "administrative_race"
    ].map(normalise_group)
    helios["genetic_ancestry"] = helios["genetic_ancestry"].map(
        normalise_group
    )
    unexpected = set(helios["genetic_ancestry"].dropna()) - set(
        GENETIC_ORDER
    )
    if unexpected:
        raise ValueError(
            f"Unexpected genetic ancestry labels: {sorted(unexpected)}"
        )

    core_raw = pd.read_csv(
        args.core,
        usecols=["FREG0_PID", "FREG5_Race", "FREG6_Race_Other"],
        low_memory=False,
    )
    core, core_duplicate_rows, core_conflicts = collapse_consistent(
        core_raw,
        "FREG0_PID",
        ["FREG5_Race", "FREG6_Race_Other"],
    )
    core["core_race"] = core["FREG5_Race"].map(normalise_group)

    interview_raw = pd.read_csv(
        args.interview,
        usecols=[
            "FREG0_PID",
            "FIAQ10_Demo1",
            "FIAQ14_1_Demo6",
            "FIAQ14_2_Demo7",
        ],
        low_memory=False,
    )
    (
        interview,
        interview_duplicate_rows,
        interview_conflicts,
    ) = collapse_consistent(
        interview_raw,
        "FREG0_PID",
        ["FIAQ10_Demo1", "FIAQ14_1_Demo6", "FIAQ14_2_Demo7"],
    )
    interview = interview.rename(
        columns={
            "FIAQ10_Demo1": "self_reported_ethnicity",
            "FIAQ14_1_Demo6": "father_ethnic_heritage",
            "FIAQ14_2_Demo7": "mother_ethnic_heritage",
        }
    )
    for column in [
        "self_reported_ethnicity",
        "father_ethnic_heritage",
        "mother_ethnic_heritage",
    ]:
        interview[column] = interview[column].map(normalise_group)

    data = helios.merge(
        core[["FREG0_PID", "core_race"]],
        left_on="participant_id",
        right_on="FREG0_PID",
        how="left",
        validate="one_to_one",
    ).drop(columns=["FREG0_PID"])
    data = data.merge(
        interview,
        left_on="participant_id",
        right_on="FREG0_PID",
        how="left",
        validate="one_to_one",
    ).drop(columns=["FREG0_PID"])
    data["parental_pair"] = data.apply(parental_pair, axis=1)

    race_overlap = (
        data["administrative_race"].notna()
        & data["core_race"].notna()
    )
    race_match = data.loc[
        race_overlap, "administrative_race"
    ].eq(data.loc[race_overlap, "core_race"])
    source_qc = pd.DataFrame(
        [
            {
                "source": args.annotation.name,
                "role": (
                    "GRID genetic ancestry, PC1-PC5, SG100K race metadata"
                ),
                "raw_rows": len(annotation_all),
                "HELIOS_rows": len(helios),
                "unique_HELIOS_participants": helios[
                    "participant_id"
                ].nunique(),
                "duplicate_identifier_rows": int(
                    annotation_all.duplicated(
                        "participant_id", keep=False
                    ).sum()
                ),
                "conflicting_identifier_n": 0,
                "sha256": sha256(args.annotation),
            },
            {
                "source": args.core.name,
                "role": "NRIC-recorded race validation",
                "raw_rows": len(core_raw),
                "HELIOS_rows": int(race_overlap.sum()),
                "unique_HELIOS_participants": int(race_overlap.sum()),
                "duplicate_identifier_rows": core_duplicate_rows,
                "conflicting_identifier_n": core_conflicts,
                "sha256": sha256(args.core),
            },
            {
                "source": args.interview.name,
                "role": "Self-described and parental ethnic heritage",
                "raw_rows": len(interview_raw),
                "HELIOS_rows": int(
                    data["self_reported_ethnicity"].notna().sum()
                ),
                "unique_HELIOS_participants": int(
                    data["self_reported_ethnicity"].notna().sum()
                ),
                "duplicate_identifier_rows": interview_duplicate_rows,
                "conflicting_identifier_n": interview_conflicts,
                "sha256": sha256(args.interview),
            },
        ]
    )
    write_tsv(source_qc, tables / "source_qc.tsv")
    write_tsv(
        variable_definitions(args.data_dictionary),
        tables / "variable_definitions.tsv",
    )
    write_tsv(
        pd.DataFrame(
            [
                {
                    "check": "Annotation race versus FREG5_Race",
                    "overlap_n": int(race_overlap.sum()),
                    "matching_n": int(race_match.sum()),
                    "agreement_percentage": (
                        100 * float(race_match.mean())
                    ),
                }
            ]
        ),
        tables / "label_lineage_checks.tsv",
    )

    genetic_counts = (
        data["genetic_ancestry"]
        .value_counts()
        .reindex(GENETIC_ORDER, fill_value=0)
        .rename_axis("genetic_ancestry")
        .reset_index(name="n")
    )
    genetic_counts["genetic_ancestry"] = genetic_counts[
        "genetic_ancestry"
    ].map(DISPLAY)
    genetic_counts["percentage"] = (
        100 * genetic_counts["n"] / len(data)
    )
    write_tsv(
        genetic_counts, tables / "genetic_ancestry_counts.tsv"
    )

    nric_cross = suppress_crosstab(
        data.dropna(
            subset=["administrative_race", "genetic_ancestry"]
        ),
        "administrative_race",
        "genetic_ancestry",
        args.minimum_cell_size,
    )
    nric_cross["administrative_race"] = nric_cross[
        "administrative_race"
    ].map(DISPLAY)
    nric_cross["genetic_ancestry"] = nric_cross[
        "genetic_ancestry"
    ].map(DISPLAY)
    write_tsv(
        nric_cross, tables / "nric_race_by_genetic_ancestry.tsv"
    )

    self_cross = suppress_crosstab(
        data.dropna(
            subset=["self_reported_ethnicity", "genetic_ancestry"]
        ),
        "self_reported_ethnicity",
        "genetic_ancestry",
        args.minimum_cell_size,
    )
    self_cross["self_reported_ethnicity"] = self_cross[
        "self_reported_ethnicity"
    ].map(DISPLAY)
    self_cross["genetic_ancestry"] = self_cross[
        "genetic_ancestry"
    ].map(DISPLAY)
    write_tsv(
        self_cross,
        tables / "selfreported_ethnicity_by_genetic_ancestry.tsv",
    )

    nric_summary = classification_summary(
        data, "administrative_race", "NRIC-recorded race"
    )
    self_summary = classification_summary(
        data,
        "self_reported_ethnicity",
        "Self-described ethnic heritage",
    )
    asymmetry_tests = statistical_tests(self_summary)
    write_tsv(
        nric_summary, tables / "nric_race_concordance_summary.tsv"
    )
    write_tsv(
        self_summary,
        tables / "selfreported_ethnicity_concordance_summary.tsv",
    )
    write_tsv(
        asymmetry_tests,
        tables / "selfreported_asymmetry_tests.tsv",
    )

    availability_rows = []
    for group in GROUPS + ["other"]:
        subset = data.loc[data["administrative_race"].eq(group)]
        available = int(
            subset["self_reported_ethnicity"].notna().sum()
        )
        low, high = wilson(available, len(subset))
        availability_rows.append(
            {
                "administrative_race": DISPLAY[group],
                "total_n": len(subset),
                "self_reported_ethnicity_recorded_n": available,
                "percentage": (
                    100 * available / len(subset)
                    if len(subset)
                    else np.nan
                ),
                "ci_low": 100 * low,
                "ci_high": 100 * high,
            }
        )
    availability = pd.DataFrame(availability_rows)
    write_tsv(
        availability,
        tables / "selfreport_availability_by_nric_race.tsv",
    )

    parent = parental_validation(data)
    write_tsv(
        parent, tables / "parental_heritage_validation.tsv"
    )
    parent_cross = suppress_crosstab(
        data.dropna(subset=["parental_pair", "genetic_ancestry"]),
        "parental_pair",
        "genetic_ancestry",
        args.minimum_cell_size,
    )
    parent_cross["genetic_ancestry"] = parent_cross[
        "genetic_ancestry"
    ].map(DISPLAY)
    write_tsv(
        parent_cross,
        tables / "parental_pair_by_genetic_ancestry.tsv",
    )

    write_tsv(
        relatedness_summary(data),
        tables / "relatedness_flag_summary.tsv",
    )
    centroid_results = centroid_sensitivity(data)
    write_tsv(
        centroid_results,
        tables / "centroid_distance_sensitivity.tsv",
    )

    if args.sg100k_covariates:
        covariates = pd.read_csv(
            args.sg100k_covariates,
            sep=r"\s+",
            dtype={"FID": str},
        )
        joined = data[["npmid", "genetic_ancestry"]].merge(
            covariates[["FID", "Sub_Cohort"]],
            left_on="npmid",
            right_on="FID",
            how="left",
            validate="one_to_one",
        )
        coverage = (
            joined.assign(
                subcohort_recorded=joined["Sub_Cohort"].notna()
            )
            .groupby("genetic_ancestry", dropna=False)[
                "subcohort_recorded"
            ]
            .agg(["sum", "count"])
            .reset_index()
            .rename(
                columns={
                    "sum": "subcohort_recorded_n",
                    "count": "total_n",
                }
            )
        )
        coverage["genetic_ancestry"] = coverage[
            "genetic_ancestry"
        ].map(DISPLAY)
        coverage["percentage"] = (
            100
            * coverage["subcohort_recorded_n"]
            / coverage["total_n"]
        )
        write_tsv(
            coverage,
            tables
            / "subcohort_coverage_by_genetic_ancestry.tsv",
        )

    group_checks = analysis_group_checks(data, args)
    if not group_checks.empty:
        write_tsv(
            group_checks, tables / "analysis_group_lineage.tsv"
        )

    scatter_plot(
        data,
        "genetic_ancestry",
        "HELIOS genetic ancestry structure",
        (
            "GRID PC1–PC2 coloured by the supplied GRID ancestry "
            "assignment"
        ),
        figures / "pca_by_genetic_ancestry.png",
        GENETIC_ORDER,
    )
    scatter_plot(
        data,
        "administrative_race",
        "HELIOS PCA by NRIC-recorded race",
        (
            "GRID PC1–PC2; race is the FREG5/SG100K "
            "administrative variable"
        ),
        figures / "pca_by_nric_race.png",
        GROUPS + ["other"],
    )
    scatter_plot(
        data,
        "self_reported_ethnicity",
        "HELIOS PCA by self-described ethnic heritage",
        (
            "GRID PC1–PC2; unrecorded questionnaire values are "
            "shown in grey"
        ),
        figures / "pca_by_selfreported_ethnicity.png",
        ["missing"] + GROUPS + ["other"],
    )
    intermediate = data.copy()
    intermediate["intermediate_highlight"] = intermediate[
        "genetic_ancestry"
    ].where(
        intermediate["genetic_ancestry"].isin(
            ["chinese-malay", "indian-chinese", "malay-indian"]
        ),
        "missing",
    )
    scatter_plot(
        intermediate,
        "intermediate_highlight",
        "Pairwise intermediate genetic ancestry",
        (
            "Core groups are grey; highlighted labels come from "
            "the GRID annotation"
        ),
        figures / "pca_intermediate_highlight.png",
        ["chinese-malay", "indian-chinese", "malay-indian"],
        highlight_only=True,
    )
    discordant = data.copy()
    discordant["discordance_highlight"] = "missing"
    for group in GROUPS:
        mask = discordant["self_reported_ethnicity"].eq(
            group
        ) & discordant["genetic_ancestry"].map(
            lambda value: group
            not in GENETIC_MEMBERS.get(value, set())
        )
        discordant.loc[mask, "discordance_highlight"] = group
    scatter_plot(
        discordant,
        "discordance_highlight",
        "Genetic labels not involving the self-described group",
        (
            "Grey points include exact and pairwise-involving "
            "assignments; highlighted points use the strict "
            "non-involving definition"
        ),
        figures / "pca_nonconcordance_highlight.png",
        GROUPS,
        highlight_only=True,
    )
    composition_plot(
        self_summary,
        "Self-described ethnic heritage and GRID genetic ancestry",
        figures
        / "selfreported_ancestry_classification_rates.png",
    )
    composition_plot(
        nric_summary,
        "NRIC-recorded race and GRID genetic ancestry",
        figures / "nric_ancestry_classification_rates.png",
    )

    report = report_markdown(
        data,
        self_summary,
        asymmetry_tests,
        parent,
        availability,
        group_checks,
        centroid_results,
    )
    write_text(
        report, root / "HELIOS_ancestry_selfreport_analysis.md"
    )

    checks = [
        {
            "check": "HELIOS participant identifiers unique",
            "status": (
                "PASS" if data["participant_id"].is_unique else "FAIL"
            ),
            "detail": f"N={len(data)}",
        },
        {
            "check": "GRID PC1-PC5 complete",
            "status": (
                "PASS"
                if data[PC_COLUMNS].notna().all().all()
                else "FAIL"
            ),
            "detail": (
                f"missing={int(data[PC_COLUMNS].isna().sum().sum())}"
            ),
        },
        {
            "check": "Annotation race matches FREG5_Race",
            "status": "PASS" if race_match.all() else "FAIL",
            "detail": (
                f"matched={int(race_match.sum())}/"
                f"{int(race_overlap.sum())}"
            ),
        },
        {
            "check": "No participant-level output written",
            "status": "PASS",
            "detail": (
                "Only aggregate TSV, report and non-identifiable "
                "figures are produced"
            ),
        },
    ]
    write_tsv(
        pd.DataFrame(checks), root / "validation_checks.tsv"
    )
    provenance = pd.DataFrame(
        [
            {"item": "python", "value": platform.python_version()},
            {"item": "numpy", "value": np.__version__},
            {"item": "pandas", "value": pd.__version__},
            {"item": "pillow", "value": Image.__version__},
            {
                "item": "minimum_cell_size",
                "value": args.minimum_cell_size,
            },
            {"item": "platform", "value": platform.platform()},
        ]
    )
    write_tsv(
        provenance, root / "software_and_settings.tsv"
    )
    if any(check["status"] != "PASS" for check in checks):
        raise RuntimeError("One or more validation checks failed")


if __name__ == "__main__":
    main()
